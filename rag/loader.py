"""Document loading for multiple file formats -> LangChain Documents."""
import csv
import json
from pathlib import Path
from typing import List

from langchain_core.documents import Document

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv", ".html", ".htm", ".pptx", ".xlsx", ".json"}


def _doc(text: str, source: str, **meta) -> Document:
    return Document(page_content=text, metadata={"source": source, **meta})


def _pdf(path: Path) -> List[Document]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return [_doc(p.extract_text() or "", path.name, page=i + 1) for i, p in enumerate(reader.pages)]


def _docx(path: Path) -> List[Document]:
    import docx2txt
    return [_doc(docx2txt.process(str(path)), path.name)]


def _text(path: Path) -> List[Document]:
    return [_doc(path.read_text(encoding="utf-8", errors="ignore"), path.name)]


def _csv(path: Path) -> List[Document]:
    with open(path, newline="", encoding="utf-8", errors="ignore") as f:
        rows = list(csv.DictReader(f))
    return [_doc("\n".join(f"{k}: {v}" for k, v in r.items()), path.name, row=i + 1)
            for i, r in enumerate(rows)]


def _html(path: Path) -> List[Document]:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return [_doc(soup.get_text(separator="\n", strip=True), path.name)]


def _pptx(path: Path) -> List[Document]:
    from pptx import Presentation
    docs = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        docs.append(_doc(text, path.name, slide=i))
    return docs


def _xlsx(path: Path) -> List[Document]:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    docs = []
    for ws in wb.worksheets:
        lines = [" | ".join("" if c is None else str(c) for c in row) for row in ws.iter_rows(values_only=True)]
        docs.append(_doc("\n".join(lines), path.name, sheet=ws.title))
    return docs


def _json(path: Path) -> List[Document]:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return [_doc(json.dumps(data, indent=2), path.name)]


LOADERS = {
    ".pdf": _pdf, ".docx": _docx, ".txt": _text, ".md": _text, ".csv": _csv,
    ".html": _html, ".htm": _html, ".pptx": _pptx, ".xlsx": _xlsx, ".json": _json,
}


def load_file(path) -> List[Document]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in LOADERS:
        raise ValueError(f"Unsupported file type: {ext}")
    return [d for d in LOADERS[ext](path) if d.page_content.strip()]


def load_directory(directory) -> List[Document]:
    docs: List[Document] = []
    for p in sorted(Path(directory).rglob("*")):
        if p.suffix.lower() in SUPPORTED:
            docs.extend(load_file(p))
    return docs
